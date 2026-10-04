"""Freeze a 5,000-item SOURCE CANDIDATE baseline, not a reviewed 5,000-item deck.
Ranking prototype deliberately exposes missing editorial components and level policy.
"""
import bisect,collections,csv,gzip,hashlib,json,math,re,xml.etree.ElementTree as ET
from pathlib import Path
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1];DATA=ROOT.parent/'woorden-research/evidence/data'
nt=list(csv.DictReader(open(DATA/'nt2lex-senses.tsv'),delimiter='\t'))
ids={r['sense_se-id'] for r in nt};od={}
with gzip.open(DATA/'odwn-1.4-checked.xml.gz','rb') as f:
 for _,e in ET.iterparse(f,events=('end',)):
  if e.tag=='LexicalEntry':
   if e.get('id') in ids:
    s=e.find('Sense');le=e.find('Lemma')
    if s is not None and le is not None:od[e.get('id')]={'definition_nl':s.get('definition'),'provenance':s.get('provenance'),'pos':e.get('partOfSpeech'),'lemma':le.get('writtenForm'),'sense_id':s.get('id'),'synset':s.get('synset')}
   e.clear()
  elif e.tag=='Synset':e.clear()
pilot=json.load(open(ROOT/'content/starter-pack.json'))['entries'];pilot_lemmas={e['lexeme']['lemma'] for e in pilot}
audit=json.load(open(ROOT/'content/legacy-audit-evidence.json'));core_lemmas={e['lemma'] for e in audit if e['pack']=='core'};optin={e['lemma'] for e in audit if e['pack']!='core'}
pool=[];seen=set()
for i,r in enumerate(nt):
 o=od.get(r['sense_se-id']);w=r['word']
 if not o or not o['definition_nl'] or 'SPEC' in r['tag'] or not w[:1].islower() or w in pilot_lemmas or w in optin:continue
 if o['lemma']!=w or o['pos'] not in ['noun','verb','adjective','adverb']:continue
 key=r['sense_se-id']
 if key in seen:continue
 seen.add(key)
 f={lv:float(r['F@'+lv]) if r['F@'+lv]!='-' else None for lv in ['A1','A2','B1','B2','C1']}
 u={lv:float(r['U@'+lv]) if r['U@'+lv]!='-' else None for lv in f}
 pool.append({'candidate_id':'nt2-odwn:'+key,'lemma':w,'pos':o['pos'],'definition_nl':o['definition_nl'],'source_ids':{'nt2lex_row':i+2,'odwn_lexical_entry':key,'odwn_sense':o['sense_id'],'odwn_synset':o['synset']},'source_lineage':o['provenance'],'nt2_f':f,'nt2_u':u,'utility':0.75 if w in core_lemmas else 0.5,'utility_basis':'legacy-editorial-core-candidate' if w in core_lemmas else 'neutral_imputation_not_reviewed','theme':'general-language-pending-review','prerequisite_value':0.5,'prerequisite_basis':'neutral_imputation_no_construction_links','situation_coverage':0.5,'situation_coverage_basis':'neutral_imputation_theme_unassigned','frequency':None})
targets={r['lemma'] for r in pool}|pilot_lemmas
freq={};wb=load_workbook(DATA/'subtlex-full.xlsx',read_only=True,data_only=True);ws=wb.active
it=ws.iter_rows(values_only=True);headers=next(it)
for row in it:
 if row[0] in targets:freq[row[0]]=dict(zip(headers,row))
wb.close();print('Frequency rows',len(freq),flush=True)
def value(v):
 try:return float(v)
 except (TypeError,ValueError):return None
map_pos={'noun':'N','verb':'WW','adjective':'ADJ','adverb':'BW'}
for r in pool:
 f=freq.get(r['lemma']);lemma_count=None;cd=None
 if f:
  cd=value(f['CDcount'])
  if f['dominant.pos.lemma']==r['lemma'] and f['dominant.pos']==map_pos.get(r['pos']):lemma_count=value(f['dominant.pos.lemma.freq'])
  r['frequency']={'surface_zipf':value(f['Zipf']),'surface_count':value(f['FREQcount']),'lemma_pos_count':lemma_count,'context_diversity_surface_proxy':cd,'source':'SUBTLEX-NL-full','row_word':f['Word'],'dominant_pos':f['dominant.pos'],'dominant_lemma':f['dominant.pos.lemma']}
 vals=[]
 r['_lc']=lemma_count;r['_cd']=cd
dist=collections.defaultdict(list)
for r in pool:
 for field in ['_lc','_cd']:
  if r[field] is not None:dist[(r['pos'],field)].append(math.log1p(r[field]))
 for lv,v in r['nt2_u'].items():
  if v is not None:dist[(r['pos'],lv)].append(math.log1p(v))
for v in dist.values():v.sort()
def pct(pos,field,v):
 if v is None:return None
 d=dist[(pos,field)];return bisect.bisect_right(d,math.log1p(v))/len(d) if d else None
levels=['A1','A2','B1','B2'];additions={'A1':600,'A2':900,'B1':1500,'B2':2000};selected=[]
for e in pilot:
 selected.append({'candidate_id':'pilot:'+e['fixture_ref'],'lemma':e['lexeme']['lemma'],'pos':e['lexeme']['pos'],'definition_nl':e['sense']['definition_nl'],'curriculum_stage':e['curriculum']['stage'],'estimated_cefr':e['cefr']['estimated'],'source_ids':{'pilot_sense_id':e['id']},'theme':e['curriculum']['themes'][0],'selection':'forced_representative_fixture','review_status':'model_draft_external_review_pending','legacy_matches':[z['legacy_id'] for z in e['legacy']]})
for band in levels:
 upto=levels[:levels.index(band)+1]
 for r in pool:
  fp=pct(r['pos'],'_lc',r['_lc']);cp=pct(r['pos'],'_cd',r['_cd'])
  gp=[pct(r['pos'],lv,r['nt2_u'][lv]) for lv in upto if r['nt2_u'][lv] is not None];gp=max(gp) if gp else None
  r['score_components']={'utility':r['utility'],'frequency':0.8*(fp if fp is not None else 0.5)+0.2*(cp if cp is not None else 0.5),'graded_exposure':gp if gp is not None else 0.5,'situation_coverage':0.5,'prerequisite_value':0.5}
  r['imputed_components']=['situation_coverage','prerequisite_value']+(['utility'] if r['utility']==0.5 else [])+(['lemma_frequency'] if fp is None else [])+(['contextual_diversity'] if cp is None else [])+(['graded_exposure'] if gp is None else [])
  r['priority_score']=sum(r['score_components'][k]*v for k,v in {'utility':35,'frequency':25,'graded_exposure':20,'situation_coverage':10,'prerequisite_value':10}.items())
  r['_has_band_evidence']=any(r['nt2_f'][lv] is not None for lv in upto)
 ranked=sorted(pool,key=lambda r:(not r['_has_band_evidence'],-r['priority_score'],r['candidate_id']))
 need=additions[band]-sum(r['curriculum_stage']==band for r in selected)
 pick=ranked[:need];picked={r['candidate_id'] for r in pick};pool=[r for r in pool if r['candidate_id'] not in picked]
 for r in pick:
  r['curriculum_stage']=band;r['estimated_cefr']=None;r['review_status']='source_candidate_not_curated';r['selection']='provisional_rank_with_unreviewed_editorial_imputations';r['legacy_matches']=[e['legacyId'] for e in audit if e['lemma']==r['lemma']];r['level_warning']='Band quota is curriculum planning, not an assessed CEFR level.'
  for k in ['_lc','_cd','_has_band_evidence']:r.pop(k,None)
  selected.append(r)
out={'version':'core-target-candidates-0.2','status':'provisional_source_candidate_baseline','count':len(selected),'cumulative_targets':{'A1':600,'A2':1500,'B1':3000,'B2':5000},'warning':'An explicit denominator for candidate gap analysis, not a claim of 5,000 fully selected, reviewed or CEFR-certified senses. Non-pilot senses require theme/utility/sense review and replacement where unsuitable.','candidates':selected}
(ROOT/'content/core-target-candidates.json').write_text(json.dumps(out,ensure_ascii=False,indent=2)+'\n')
with (ROOT/'content/core-missing-candidates.csv').open('w',newline='',encoding='utf-8-sig') as f:
 wr=csv.DictWriter(f,fieldnames=['candidate_id','lemma','pos','definition_nl','curriculum_stage','review_status']);wr.writeheader();wr.writerows({k:r[k] for k in wr.fieldnames} for r in selected if not r['legacy_matches'])
summary={'baseline':out['version'],'candidate_count':len(selected),'distinct_lemmas':len({r['lemma'] for r in selected}),'headword_matched_candidates':sum(bool(r['legacy_matches']) for r in selected),'headword_missing_candidates':sum(not r['legacy_matches'] for r in selected),'confirmed_sense_coverage':None,'limitation':'A headword match is an upper-bound candidate signal, not verified sense coverage. No full 5,000-sense coverage percentage is claimed.','stages':dict(collections.Counter(r['curriculum_stage'] for r in selected)),'pilot_missing_lemmas':[{'ref':e['fixture_ref'],'lemma':e['lexeme']['lemma']} for e in pilot if not e['legacy']],'seed_lemmas_outside_baseline':len({e['lemma'] for e in audit}-{r['lemma'] for r in selected})}
(ROOT/'evidence/core-coverage-summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(summary,ensure_ascii=False,indent=2))
